"""
Proposal Submission Pipeline - team tracker (Streamlit + Supabase/Postgres)

Run:              streamlit run app.py
Requirements:     see requirements.txt

Storage
  * If a `database_url` secret is set (Streamlit Cloud -> App settings -> Secrets),
    data lives in that Postgres database (Supabase). It survives restarts.
  * If not, the app falls back to a local SQLite file (proposals.db) so you can
    still run and test it on your own computer.

Everyone who can reach the app can VIEW the tracker. Only people who enter the
shared password can add / edit / delete proposals.
"""
from __future__ import annotations

import hmac
import io
import os
import sqlite3
import time
import uuid
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st

# ─────────────────────────────────────────────────────────────────────────────
# Settings - edit these lists to add a person / status / submission mode
# ─────────────────────────────────────────────────────────────────────────────
APP_TITLE = "Proposal Submission Pipeline"
APP_SUBTITLE = "Deadlines, owners and submission status for every proposal in one place"
TIMEZONE = "Asia/Karachi"          # used for "today" / overdue calculations
SOON_DAYS = 3                      # deadlines within this many days count as "soon"

PEOPLE = [
    "Irfan", "Ali Waheed", "Sajida", "Ayesha", "Zunaira",
    "Shahid Minhas", "Maryam", "Amina", "Sajjad Haider", "Daniyal", "Rana Nazir"
]
MODES = [
    "Email only",
    "Email + Hard copy",
    "Online/Portal only",
    "Online/Portal + Hard copy",
    "Hard copy only",
    "Email + Online/Portal",
    "Email + Online/Portal + Hard copy",
]
STATUSES = ["Not Started Yet", "In Progress", "Submitted", "Not Submitted", "Declined", "Accepted"]
CLOSED_STATUSES = {"Submitted", "Not Submitted", "Declined", "Accepted"}   # no longer "open" work

# Column order here == the SELECT order in load_data()
COLUMNS = ["ID", "Project No", "Proposal", "Deadline", "Submission Mode",
           "Assigned to", "Status", "Remarks"]
DB_COLUMNS = ["id", "project_no", "proposal", "deadline", "submission_mode",
              "assigned_to", "status", "remarks"]
TEXT_COLUMNS = ["ID", "Proposal", "Submission Mode", "Assigned to", "Status", "Remarks"]

# Local SQLite file: used when no database_url is set, and as a one-time import
# source for Supabase (see _import_legacy_sqlite).
DB_FILE = Path(os.environ.get("PROPOSALS_DB") or Path(__file__).with_name("proposals.db"))
# If present, imported once when a brand-new local database is created.
SEED_FILE = Path(__file__).with_name("seed_proposals.csv")

# Row colour by Status. Semi-transparent so they look right in both light and dark themes.
# A status that isn't listed here (e.g. "Accepted") gets no colour.
ROW_COLORS = {
    "Submitted": "rgba(34, 197, 94, 0.16)",       # green
    "In Progress": "rgba(245, 158, 11, 0.24)",    # amber
    "Not Started Yet": "rgba(56, 189, 248, 0.22)",  # light blue
    "Declined": "rgba(220, 38, 38, 0.20)",        # red
    "Not Submitted": "rgba(220, 38, 38, 0.20)",   # red
}
# Same colours for the Excel export (hex, no #).
EXCEL_FILLS = {
    "Submitted": "D5F0DC",
    "In Progress": "FCE5B0",
    "Not Started Yet": "D6EEFB",
    "Declined": "F8CFCF",
    "Not Submitted": "F8CFCF",
}
# Summary-tile accent colours
TILE_COLORS = {
    "overdue": "#ef4444",   # red
    "soon": "#6366f1",      # indigo
    "progress": "#f59e0b",  # amber
    "submitted": "#22c55e", # green
}


# ─────────────────────────────────────────────────────────────────────────────
# Small helpers
# ─────────────────────────────────────────────────────────────────────────────
def today_local() -> date:
    try:
        return datetime.now(ZoneInfo(TIMEZONE)).date()
    except Exception:  # tzdata missing - fall back to the machine's date
        return date.today()


def _secret(name: str, default=None):
    """Read from Streamlit secrets, falling back to an environment variable."""    
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.environ.get(name.upper(), default)


def _width_kwargs() -> dict:
    """st.dataframe 'stretch' syntax differs between Streamlit versions."""
    try:
        major, minor = (int(x) for x in st.__version__.split(".")[:2])
        if (major, minor) >= (1, 50):
            return {"width": "stretch"}
    except Exception:
        pass
    return {"use_container_width": True}


def new_id() -> str:
    return "p" + uuid.uuid4().hex[:8]     # internal row key (not the visible Project ID)


# ─────────────────────────────────────────────────────────────────────────────
# Storage  (Postgres/Supabase when database_url is set, otherwise SQLite)
# All SQL below is written with %s placeholders and runs on both databases.
# ─────────────────────────────────────────────────────────────────────────────
def normalise(raw: pd.DataFrame) -> pd.DataFrame:
    """Coerce whatever came back from storage into a clean, typed frame."""
    df = raw.copy() if raw is not None else pd.DataFrame()
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""
    df = df[COLUMNS].dropna(how="all")
    for col in TEXT_COLUMNS:
        df[col] = df[col].fillna("").astype(str).str.strip()
    df["Project No"] = pd.to_numeric(df["Project No"], errors="coerce").astype("Int64")
    df["Deadline"] = pd.to_datetime(df["Deadline"], errors="coerce", format="mixed")
    df = df[df["Proposal"] != ""].copy()
    missing = df["ID"] == ""
    df.loc[missing, "ID"] = [new_id() for _ in range(int(missing.sum()))]
    # Visible Project ID: 001, 002, ...
    df["Project ID"] = df["Project No"].map(lambda n: f"{int(n):03d}" if pd.notna(n) else "")
    return df.reset_index(drop=True)


def _database_url() -> str:
    url = str(_secret("database_url") or "").strip()
    return url if url.startswith(("postgres://", "postgresql://")) else ""


def use_postgres() -> bool:
    return bool(_database_url())


def _connect():
    if use_postgres():
        import psycopg2

        return psycopg2.connect(_database_url(), connect_timeout=10)
    con = sqlite3.connect(DB_FILE, timeout=15)   # wait if a teammate is mid-write
    con.execute("PRAGMA journal_mode=WAL")       # readers don't block the writer
    return con


def _sql(sql: str) -> str:
    return sql if use_postgres() else sql.replace("%s", "?")


def _write(sql: str, params: tuple = ()) -> None:
    with closing(_connect()) as con:
        cur = con.cursor()
        cur.execute(_sql(sql), params)
        con.commit()


def _query(sql: str, params: tuple = ()) -> list:
    with closing(_connect()) as con:
        cur = con.cursor()
        cur.execute(_sql(sql), params)
        return cur.fetchall()


_DDL_PROPOSALS = """CREATE TABLE IF NOT EXISTS proposals (
    id TEXT PRIMARY KEY,
    project_no INTEGER UNIQUE,
    proposal TEXT NOT NULL,
    deadline TEXT,
    submission_mode TEXT,
    assigned_to TEXT,
    status TEXT,
    remarks TEXT
)"""
_DDL_META = "CREATE TABLE IF NOT EXISTS app_meta (key TEXT PRIMARY KEY, value TEXT)"


def _harden_postgres(cur) -> None:
    """Supabase exposes tables to its public web API unless row-level security is on.
    With RLS enabled and no policies, only our direct database connection can read/write."""
    for table in ("proposals", "app_meta"):
        cur.execute(f"ALTER TABLE {table} ENABLE ROW LEVEL SECURITY")


def _upgrade_old_sqlite(cur) -> None:
    """Add the Project ID / Remarks columns to a pre-existing local proposals.db."""
    have = {row[1] for row in cur.execute("PRAGMA table_info(proposals)").fetchall()}
    if "project_no" not in have:
        cur.execute("ALTER TABLE proposals ADD COLUMN project_no INTEGER")
    if "remarks" not in have:
        cur.execute("ALTER TABLE proposals ADD COLUMN remarks TEXT")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS proposals_project_no ON proposals (project_no)")


def _import_legacy_sqlite() -> None:
    """One-time: copy rows from a proposals.db sitting next to the app into Postgres.
    Runs only if the database is still empty, and never twice (marker in app_meta)."""
    if not DB_FILE.exists():
        return
    if _query("SELECT 1 FROM app_meta WHERE key = %s", ("legacy_import",)):
        return
    if _query("SELECT COUNT(*) FROM proposals")[0][0] == 0:
        try:
            with closing(sqlite3.connect(DB_FILE.resolve().as_uri() + "?mode=ro", uri=True)) as old:
                have = [row[1] for row in old.execute("PRAGMA table_info(proposals)")]
                select = ", ".join(c if c in have else "NULL" for c in DB_COLUMNS)
                rows = [list(r) for r in old.execute(f"SELECT {select} FROM proposals ORDER BY rowid")]
        except sqlite3.Error:
            return                      # not a usable proposals.db - leave it alone
        # Keep existing Project Nos; number the rest in the order they were created.
        nxt = max((r[1] for r in rows if r[1] is not None), default=0) + 1
        for r in rows:
            if r[1] is None:
                r[1], nxt = nxt, nxt + 1
        with closing(_connect()) as con:
            cur = con.cursor()
            for r in rows:
                cur.execute(
                    _sql(
                        "INSERT INTO proposals (id, project_no, proposal, deadline, submission_mode, "
                        "assigned_to, status, remarks) VALUES (%s, %s, %s, %s, %s, %s, %s, %s) "
                        "ON CONFLICT (id) DO NOTHING"
                    ),
                    tuple(r),
                )
            con.commit()
    _write(
        "INSERT INTO app_meta (key, value) VALUES (%s, %s) ON CONFLICT (key) DO NOTHING",
        ("legacy_import", "done"),
    )


def _backfill_project_numbers() -> None:
    """Give a Project No to any row that lacks one (e.g. added by hand in the database)."""
    order = "id" if use_postgres() else "rowid"
    for (pid,) in _query(f"SELECT id FROM proposals WHERE project_no IS NULL ORDER BY {order}"):
        _write(
            "UPDATE proposals SET project_no = "
            "(SELECT COALESCE(MAX(project_no), 0) + 1 FROM proposals) WHERE id = %s",
            (pid,),
        )


def init_db() -> None:
    first_local_run = not use_postgres() and not DB_FILE.exists()
    with closing(_connect()) as con:
        cur = con.cursor()
        cur.execute(_DDL_PROPOSALS)
        cur.execute(_DDL_META)
        if use_postgres():
            _harden_postgres(cur)
        else:
            _upgrade_old_sqlite(cur)
        con.commit()
    if use_postgres():
        _import_legacy_sqlite()
    elif first_local_run and SEED_FILE.exists():
        for rec in normalise(pd.read_csv(SEED_FILE, dtype=str)).to_dict("records"):
            upsert_proposal(rec)
    _backfill_project_numbers()


@st.cache_resource(show_spinner=False)
def _init_once(_url: str) -> bool:
    """Create tables / import old data once per server start, not on every page refresh."""
    init_db()
    return True


def load_data() -> pd.DataFrame:
    _init_once(_database_url())
    rows = _query(
        "SELECT id, project_no, proposal, deadline, submission_mode, assigned_to, status, remarks "
        "FROM proposals"
    )
    return normalise(pd.DataFrame(rows, columns=COLUMNS))


def upsert_proposal(record: dict) -> None:
    """Insert, or update if the ID already exists. Only touches that one row.
    New rows get the next Project No automatically; existing rows keep theirs."""
    deadline = pd.Timestamp(record["Deadline"])
    deadline = "" if pd.isna(deadline) else deadline.strftime("%Y-%m-%d")
    params = (
        record.get("ID") or new_id(), record["Proposal"], deadline,
        record["Submission Mode"], record["Assigned to"], record["Status"],
        record.get("Remarks", ""),
    )
    sql = (
        "INSERT INTO proposals (id, project_no, proposal, deadline, submission_mode, "
        "assigned_to, status, remarks) "
        "VALUES (%s, (SELECT COALESCE(MAX(project_no), 0) + 1 FROM proposals), %s, %s, %s, %s, %s, %s) "
        "ON CONFLICT (id) DO UPDATE SET "
        "proposal = excluded.proposal, deadline = excluded.deadline, "
        "submission_mode = excluded.submission_mode, assigned_to = excluded.assigned_to, "
        "status = excluded.status, remarks = excluded.remarks"
    )
    for attempt in range(5):
        try:
            _write(sql, params)
            return
        except Exception as exc:
            # Two people adding at the same instant can pick the same number: try again.
            clash = type(exc).__name__ in ("IntegrityError", "UniqueViolation")
            if not clash or attempt == 4:
                raise
            time.sleep(0.05)


def delete_proposal(pid: str) -> None:
    _write("DELETE FROM proposals WHERE id = %s", (pid,))


# ─────────────────────────────────────────────────────────────────────────────
# Presentation logic
# ─────────────────────────────────────────────────────────────────────────────
def decorate(df: pd.DataFrame, today: date) -> pd.DataFrame:
    """Add 'Time left' text and an internal urgency key, and sort open work first."""
    d = df.copy()
    days = (d["Deadline"].dt.normalize() - pd.Timestamp(today)).dt.days

    labels, urgency = [], []
    for status, n in zip(d["Status"], days):
        if status == "Submitted":
            labels.append("")
            urgency.append("submitted")
        elif status == "Not Submitted":
            labels.append("")
            urgency.append("missed")
        elif pd.isna(n):
            labels.append("")
            urgency.append("ok")
        elif n < 0:
            labels.append(f"{-int(n)} day{'s' if n != -1 else ''} overdue")
            urgency.append("overdue")
        elif n == 0:
            labels.append("Due today")
            urgency.append("soon")
        elif n == 1:
            labels.append("Tomorrow")
            urgency.append("soon")
        else:
            labels.append(f"{int(n)} days")
            urgency.append("soon" if n <= SOON_DAYS else "ok")

    d["Time left"] = labels
    d["_urgency"] = urgency
    d["_days"] = days

    closed = d["Status"].isin(CLOSED_STATUSES)
    open_rows = d[~closed].sort_values("Deadline", na_position="last")
    closed_rows = d[closed].sort_values("Deadline", ascending=False, na_position="last")
    return pd.concat([open_rows, closed_rows], ignore_index=True)


def split_people(value: str) -> list[str]:
    return [p.strip() for p in str(value).split(",") if p.strip()]


def apply_filters(d: pd.DataFrame, search, statuses, people, open_only, window) -> pd.DataFrame:
    out = d
    if search and search.strip():
        s = search.strip()
        out = out[
            out["Proposal"].str.contains(s, case=False, regex=False)
            | out["Project ID"].str.contains(s, case=False, regex=False)
            | out["Remarks"].str.contains(s, case=False, regex=False)
        ]
    if statuses:
        out = out[out["Status"].isin(statuses)]
    if people:
        out = out[out["Assigned to"].map(lambda v: any(p in split_people(v) for p in people))]
    if open_only:
        out = out[~out["Status"].isin(CLOSED_STATUSES)]
    if window == "Overdue":
        out = out[out["_urgency"] == "overdue"]
    elif window == "Next 7 days":
        out = out[(~out["Status"].isin(CLOSED_STATUSES)) & out["_days"].between(0, 7)]
    elif window == "This month":
        t = pd.Timestamp(today_local())
        out = out[(out["Deadline"].dt.year == t.year) & (out["Deadline"].dt.month == t.month)]
    return out.reset_index(drop=True)


def to_excel_bytes(d: pd.DataFrame) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = ["Project ID", "Proposal", "Deadline", "Submission Mode", "Assigned to", "Status", "Remarks"]
    out = d[cols].copy()
    out["Deadline"] = out["Deadline"].dt.date

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        out.to_excel(xw, sheet_name="Proposal Pipeline", index=False)
        ws = xw.sheets["Proposal Pipeline"]

        head_fill = PatternFill("solid", start_color="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        widths = {"A": 11, "B": 52, "C": 12, "D": 32, "E": 34, "F": 16, "G": 60}
        for letter, w in widths.items():
            ws.column_dimensions[letter].width = w
        for i, status in enumerate(d["Status"], start=2):
            ws.cell(row=i, column=3).number_format = "d-mmm-yy"
            ws.cell(row=i, column=7).alignment = Alignment(wrap_text=True, vertical="top")
            fill = EXCEL_FILLS.get(status)
            if fill:
                for c in range(1, len(cols) + 1):
                    ws.cell(row=i, column=c).fill = PatternFill("solid", start_color=fill)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(out) + 1, 2)}"
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Styling (page header, summary tiles, section headings)
# Colours are semi-transparent / inherited so they work in light and dark themes.
# ─────────────────────────────────────────────────────────────────────────────
_CSS = """
[data-testid="stMainBlockContainer"], .block-container { padding-top: 3.4rem; padding-bottom: 2.5rem; }
.app-banner {
  display: flex; justify-content: space-between; align-items: center; gap: 1rem; flex-wrap: wrap;
  padding: 1.1rem 1.4rem; border-radius: 14px;
  border: 1px solid rgba(148,163,184,.35); border-left: 6px solid #3b82f6;
  background: linear-gradient(135deg, rgba(59,130,246,.16), rgba(99,102,241,.08));
}
.app-title { font-size: 2rem; font-weight: 800; letter-spacing: -.02em; line-height: 1.15; }
.app-sub { margin-top: .3rem; font-size: .95rem; opacity: .72; }
.app-date { text-align: right; font-size: .82rem; opacity: .8; line-height: 1.35; }
.app-date strong { font-size: 1rem; }
.sec-head {
  display: flex; align-items: center; gap: .7rem; margin: 1.5rem 0 .75rem;
  font-size: .78rem; font-weight: 700; letter-spacing: .14em; text-transform: uppercase; opacity: .8;
}
.sec-head::after { content: ""; flex: 1; height: 1px; background: rgba(148,163,184,.4); }
.tiles { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: .9rem; }
@media (max-width: 760px) { .tiles { grid-template-columns: repeat(2, minmax(0, 1fr)); } }
.tile {
  padding: .8rem 1.05rem .9rem; border-radius: 12px;
  border: 1px solid rgba(148,163,184,.35); border-top: 4px solid var(--c);
  background: rgba(148,163,184,.07);
}
.tile .lbl { font-size: .75rem; font-weight: 600; letter-spacing: .09em; text-transform: uppercase; opacity: .72; }
.tile .val { margin-top: .15rem; font-size: 2.3rem; font-weight: 800; line-height: 1.1; color: var(--c); }
"""


def _inject_css() -> None:
    st.markdown("<style>" + _CSS + "</style>", unsafe_allow_html=True)


def _banner_html(today: date) -> str:
    return (
        '<div class="app-banner"><div>'
        f'<div class="app-title">{APP_TITLE}</div>'
        f'<div class="app-sub">{APP_SUBTITLE}</div></div>'
        f'<div class="app-date">{today:%A}<br><strong>{today:%d %b %Y}</strong></div></div>'
    )


def _section_html(label: str) -> str:
    return f'<div class="sec-head">{label}</div>'


def _tiles_html(items: list[tuple[str, int, str]]) -> str:
    cards = "".join(
        f'<div class="tile" style="--c:{color}"><div class="lbl">{label}</div>'
        f'<div class="val">{int(value)}</div></div>'
        for label, value, color in items
    )
    return f'<div class="tiles">{cards}</div>'


# ─────────────────────────────────────────────────────────────────────────────
# Access control
# ─────────────────────────────────────────────────────────────────────────────
def is_editor() -> bool:
    return bool(st.session_state.get("is_editor"))


def password_ok(attempt: str) -> bool:
    expected = _secret("edit_password")
    if not expected:
        return False
    return hmac.compare_digest(attempt.encode(), str(expected).encode())


def sidebar_access() -> None:
    with st.sidebar:
        st.subheader("Access")
        if is_editor():
            st.success("Editing unlocked")
            if st.button("Lock"):
                st.session_state["is_editor"] = False
                st.rerun()
            return

        st.caption("View-only. Enter the team password to add or edit proposals.")
        if not _secret("edit_password"):
            st.warning("No `edit_password` secret is set, so editing is disabled.")
            return
        with st.form("unlock"):
            attempt = st.text_input("Password", type="password")
            go = st.form_submit_button("Unlock editing")
        if go:
            if password_ok(attempt):
                st.session_state["is_editor"] = True
                st.rerun()
            else:
                time.sleep(1)      # slow down guessing
                st.error("Wrong password.")


# ─────────────────────────────────────────────────────────────────────────────
# Dialogs (add / edit / delete)
# ─────────────────────────────────────────────────────────────────────────────
def _options(base: list[str], current: list[str]) -> list[str]:
    """Keep legacy values (e.g. someone no longer in PEOPLE) selectable."""
    return base + [c for c in current if c and c not in base]


def _bump_table() -> None:
    st.session_state["table_version"] = st.session_state.get("table_version", 0) + 1


def _proposal_form(record: dict | None, form_key: str) -> None:
    rec = record or {}
    current_people = split_people(rec.get("Assigned to", ""))
    deadline_default = rec.get("Deadline")
    if deadline_default is None or pd.isna(deadline_default):
        deadline_default = today_local()
    deadline_default = pd.Timestamp(deadline_default).date()

    mode_opts = _options(MODES, [rec.get("Submission Mode", "")])
    status_opts = _options(STATUSES, [rec.get("Status", "")])
    mode_idx = mode_opts.index(rec["Submission Mode"]) if rec.get("Submission Mode") in mode_opts else None
    status_val = rec.get("Status") or "Not Started Yet"

    with st.form(form_key):
        if rec.get("Project ID"):
            st.caption(f"Project ID: **{rec['Project ID']}**")
        else:
            st.caption("The Project ID is assigned automatically when you save.")
        name = st.text_input("Proposal", value=rec.get("Proposal", ""))
        deadline = st.date_input("Deadline", value=deadline_default, format="DD/MM/YYYY")
        mode = st.selectbox("Submission mode", mode_opts, index=mode_idx, placeholder="Choose a mode")
        people = st.multiselect(
            "Assigned to", _options(PEOPLE, current_people), default=current_people,
            placeholder="Pick one or more people",
        )
        status = st.selectbox("Status", status_opts, index=status_opts.index(status_val))
        remarks = st.text_area("Remarks", value=rec.get("Remarks", ""), height=100,
                               placeholder="Optional notes, links, next steps…")
        saved = st.form_submit_button("Save", type="primary")

    if not saved:
        return
    if not name.strip():
        st.error("Please enter the proposal name.")
        return
    if not mode:
        st.error("Please choose a submission mode.")
        return
    try:
        upsert_proposal({
            "ID": rec.get("ID", ""),
            "Proposal": name.strip(),
            "Deadline": deadline,
            "Submission Mode": mode,
            "Assigned to": ", ".join(people),
            "Status": status,
            "Remarks": remarks.strip(),
        })
    except Exception as exc:
        st.error(f"Could not save: {exc}")
        return
    _bump_table()
    st.rerun()


@st.dialog("Add proposal")
def add_dialog() -> None:
    _proposal_form(None, "add_form")


@st.dialog("Edit proposal")
def edit_dialog(record: dict) -> None:
    _proposal_form(record, "edit_form")


@st.dialog("Delete proposal")
def delete_dialog(pid: str, name: str) -> None:
    st.write(f"Delete **{name}**? This can't be undone.")
    c1, c2 = st.columns(2)
    if c1.button("Yes, delete", type="primary"):
        try:
            delete_proposal(pid)
        except Exception as exc:
            st.error(f"Could not delete: {exc}")
            return
        _bump_table()
        st.rerun()
    if c2.button("Cancel"):
        st.rerun()


# ─────────────────────────────────────────────────────────────────────────────
# Page
# ─────────────────────────────────────────────────────────────────────────────
def main() -> None:
    st.set_page_config(
        page_title=APP_TITLE, page_icon="📋", layout="wide",
        initial_sidebar_state="collapsed",       # sidebar opens only when the user opens it
    )
    _inject_css()
    sidebar_access()

    today = today_local()
    st.markdown(_banner_html(today), unsafe_allow_html=True)

    try:
        df = load_data()
    except Exception as exc:
        where = "the Supabase database" if use_postgres() else f"the local database ({DB_FILE})"
        st.error(f"Couldn't connect to {where}: {str(exc).splitlines()[0] if str(exc) else type(exc).__name__}")
        st.stop()

    data = decorate(df, today)
    open_mask = ~data["Status"].isin(CLOSED_STATUSES)

    # Summary tiles
    st.markdown(
        _section_html("Overview")
        + _tiles_html([
            ("Overdue", int((data["_urgency"] == "overdue").sum()), TILE_COLORS["overdue"]),
            ("Due in next 7 days", int((open_mask & data["_days"].between(0, 7)).sum()), TILE_COLORS["soon"]),
            ("In progress", int((data["Status"] == "In Progress").sum()), TILE_COLORS["progress"]),
            ("Submitted", int((data["Status"] == "Submitted").sum()), TILE_COLORS["submitted"]),
        ]),
        unsafe_allow_html=True,
    )

    # Filters
    st.markdown(_section_html("Proposals"), unsafe_allow_html=True)
    with st.container(border=True):
        f1, f2, f3, f4, f5 = st.columns([2.4, 1.5, 1.5, 1.4, 1.5])
        search = f1.text_input("Search", placeholder="Name, Project ID or remarks…")
        statuses = f2.multiselect("Status", _options(STATUSES, sorted(set(data["Status"]) - set(STATUSES))))
        all_people = _options(PEOPLE, sorted({p for v in data["Assigned to"] for p in split_people(v)}))
        people = f3.multiselect("Assigned to", all_people)
        window = f4.selectbox("Deadline", ["Any time", "Overdue", "Next 7 days", "This month"])
        open_only = f5.toggle("Open only", value=False, help="Hide Submitted / Not Submitted / Declined / Accepted")

    view = apply_filters(data, search, statuses, people, open_only, window)

    # Table
    display_cols = ["Project ID", "Proposal", "Deadline", "Time left", "Submission Mode",
                    "Assigned to", "Status", "Remarks"]
    styled = view.style.apply(
        lambda row: [f"background-color: {ROW_COLORS.get(row['Status'], '')}"] * len(row), axis=1
    )
    event = st.dataframe(
        styled,
        hide_index=True,
        column_order=display_cols,
        column_config={
            "Project ID": st.column_config.TextColumn("Project ID", width="small"),
            "Proposal": st.column_config.TextColumn("Proposal", width="large"),
            "Deadline": st.column_config.DateColumn("Deadline", format="DD MMM YYYY"),
            "Time left": st.column_config.TextColumn("Time left", width="small"),
            "Remarks": st.column_config.TextColumn("Remarks", width="large"),
        },
        height=max(120, min(36 * (len(view) + 1) + 4, 720)),
        on_select="rerun" if is_editor() else "ignore",
        selection_mode="single-row",
        key=f"table_{st.session_state.get('table_version', 0)}",
        **_width_kwargs(),
    )
    st.caption(
        f"{len(view)} of {len(data)} proposals · green = submitted · amber = in progress "
        "· light blue = not started yet · red = declined / not submitted"
    )

    # Actions
    selected = None
    if is_editor():
        rows = event.selection.rows if event is not None else []
        if rows and rows[0] < len(view):
            selected = view.iloc[rows[0]].to_dict()

    a1, a2, a3, _, a5 = st.columns([1.1, 1.1, 1.1, 2.2, 1.3])
    if is_editor():
        if a1.button("➕ Add proposal", type="primary"):
            add_dialog()
        if a2.button("✏️ Edit selected", disabled=selected is None):
            edit_dialog(selected)
        if a3.button("🗑️ Delete", disabled=selected is None):
            delete_dialog(selected["ID"], selected["Proposal"])
        if selected is None:
            st.caption("Tip: click a row's checkbox (left edge of the table) to edit or delete it.")
    else:
        st.caption("🔒 View-only — to edit, open the sidebar (arrow at the top-left) and enter the team password.")

    a5.download_button(
        "⬇️ Export to Excel",
        data=to_excel_bytes(view),
        file_name=f"Proposal_Submission_Pipeline_{today:%Y-%m-%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help="Exports the rows currently shown (respecting your filters).",
    )


if __name__ == "__main__":
    main()
